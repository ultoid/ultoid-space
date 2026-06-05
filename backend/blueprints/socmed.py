import os
import json
import re
import csv
from datetime import datetime
from flask import Blueprint, request, jsonify
import requests
from bs4 import BeautifulSoup
from apify_client import ApifyClient
from ..utils import login_required
from ..config import IG_ACCOUNTS_FILE, TT_ACCOUNTS_FILE, FB_ACCOUNTS_FILE, GMAPS_ACCOUNTS_FILE, GMAPS_REVIEWS_FILE

socmed_bp = Blueprint('socmed', __name__)


def load_json_file(path):
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}


def save_json_file(path, data):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)


@socmed_bp.route('/api/ig/accounts', methods=['GET', 'POST', 'DELETE'])
@login_required
def manage_ig_accounts():
    accounts = load_json_file(IG_ACCOUNTS_FILE)
    if request.method == 'GET':
        last_updated = '-'
        if os.path.exists(IG_ACCOUNTS_FILE):
            last_updated = datetime.fromtimestamp(os.path.getmtime(IG_ACCOUNTS_FILE)).strftime('%d %b %Y, %H:%M:%S')
        return jsonify({'accounts': accounts, 'last_updated': last_updated})
    if request.method == 'POST':
        data = request.json
        username = data.get('profile', {}).get('username')
        if username:
            accounts[username] = data
            save_json_file(IG_ACCOUNTS_FILE, accounts)
        return jsonify({'success': True})
    username = request.json.get('username')
    if username in accounts:
        accounts.pop(username, None)
        save_json_file(IG_ACCOUNTS_FILE, accounts)
    return jsonify({'success': True})


@socmed_bp.route('/api/scrape/ig', methods=['POST'])
@login_required
def api_scrape_ig():
    data = request.json
    username = data.get('username', '').replace('@', '').strip().lower()
    if not username:
        return jsonify({'success': False, 'error': 'Username kosong'})
    try:
        url_endpoint = 'https://instagram120.p.rapidapi.com/api/instagram/profile'
        _rapid_key = os.environ.get('RAPIDAPI_KEY', '')
        if not _rapid_key:
            return jsonify({'success': False, 'error': 'RAPIDAPI_KEY belum diset di environment variable server.'})
        headers_rapid = {
            'content-type': 'application/json',
            'X-RapidAPI-Key': _rapid_key,
            'X-RapidAPI-Host': 'instagram120.p.rapidapi.com'
        }
        res_rapid = requests.post(url_endpoint, json={'username': username}, headers=headers_rapid, timeout=12)
        if res_rapid.status_code == 200:
            json_data = res_rapid.json()
            result_obj = json_data.get('result', {})
            if result_obj:
                followers_num = result_obj.get('edge_followed_by', {}).get('count', 0)
                posts_num = result_obj.get('edge_owner_to_timeline_media', {}).get('count', 0)
                bio_text = result_obj.get('biography', '-')
                full_name = result_obj.get('full_name', username)
                profile_pic = result_obj.get('profile_pic_url', f'https://ui-avatars.com/api/?name={username}&background=random')
                return jsonify({
                    'success': True,
                    'is_fallback': False,
                    'profile': {
                        'username': username,
                        'profile_pic': profile_pic,
                        'bio': f'{full_name} | {bio_text}',
                        'followers': followers_num,
                        'total_posts': posts_num
                    },
                    'posts': []
                })
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"RapidAPI failed: {e}")
    try:
        followers_num = 0
        posts_num = 0
        bio_name = username
        profile_pic = f'https://ui-avatars.com/api/?name={username}&background=random'
        html_url = f'https://www.instagram.com/{username}/'
        res_html = requests.get(html_url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}, timeout=10)
        soup = BeautifulSoup(res_html.text, 'html.parser')
        meta_desc = soup.find('meta', attrs={'name': 'description'}) or soup.find('meta', property='og:description')
        if meta_desc:
            content = meta_desc.get('content', '')
            f_match = re.search(r'([\d,KMB.]+)\s+Followers', content, re.IGNORECASE)
            if f_match:
                f_str = f_match.group(1).upper().replace(',', '')
                if 'K' in f_str:
                    followers_num = int(float(f_str.replace('K', '')) * 1000)
                elif 'M' in f_str:
                    followers_num = int(float(f_str.replace('M', '')) * 1000000)
                else:
                    followers_num = int(float(f_str))
            p_match = re.search(r'([\d,KMB.]+)\s+Posts', content, re.IGNORECASE)
            if p_match:
                p_str = p_match.group(1).upper().replace(',', '')
                if 'K' in p_str:
                    posts_num = int(float(p_str.replace('K', '')) * 1000)
                else:
                    posts_num = int(float(p_str))
        if soup.title:
            title_match = re.search(r'^(.*?)\s*\(@', soup.title.string)
            if title_match:
                bio_name = title_match.group(1).strip()
        return jsonify({
            'success': True,
            'is_fallback': True,
            'profile': {
                'username': username,
                'profile_pic': profile_pic,
                'bio': f'{bio_name}',
                'followers': followers_num,
                'total_posts': posts_num
            },
            'posts': []
        })
    except Exception:
        return jsonify({'success': False, 'error': 'Koneksi terputus atau gagal memproses halaman.'})


@socmed_bp.route('/api/ig/auto_backup', methods=['POST'])
@login_required
def ig_auto_backup():
    accounts = load_json_file(IG_ACCOUNTS_FILE)
    if not accounts:
        return jsonify({'success': False, 'error': 'Database Kosong'})
    backup_dir = os.path.join('static', 'file_manager', 'Social Media Report', 'IG Report')
    os.makedirs(backup_dir, exist_ok=True)
    filename = f'IG_Report_AutoBackup_{datetime.now().strftime("%Y-%m-%d_%H%M")}.csv'
    filepath = os.path.join(backup_dir, filename)
    usernames = sorted(accounts.keys(), key=lambda x: accounts[x].get('profile', {}).get('followers', 0), reverse=True)
    with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(['No', 'Username', 'Nama Profil', 'Followers', 'Posts'])
        for i, u in enumerate(usernames):
            prof = accounts[u].get('profile', {})
            writer.writerow([
                i + 1,
                f"@{prof.get('username', u)}",
                prof.get('bio', '-').replace('\n', ' '),
                prof.get('followers', 0),
                prof.get('total_posts', 0)
            ])
    return jsonify({'success': True, 'filename': filename})


@socmed_bp.route('/api/tt/accounts', methods=['GET', 'POST', 'DELETE'])
@login_required
def manage_tt_accounts():
    accounts = load_json_file(TT_ACCOUNTS_FILE)
    if request.method == 'GET':
        last_updated = '-'
        if os.path.exists(TT_ACCOUNTS_FILE):
            last_updated = datetime.fromtimestamp(os.path.getmtime(TT_ACCOUNTS_FILE)).strftime('%d %b %Y, %H:%M:%S')
        return jsonify({'accounts': accounts, 'last_updated': last_updated})
    if request.method == 'POST':
        data = request.json
        key_id = data.get('profile', {}).get('user_id') or data.get('profile', {}).get('username')
        if key_id:
            accounts[key_id] = data
            save_json_file(TT_ACCOUNTS_FILE, accounts)
        return jsonify({'success': True})
    key_id = request.json.get('username')
    if key_id in accounts:
        accounts.pop(key_id, None)
        save_json_file(TT_ACCOUNTS_FILE, accounts)
    return jsonify({'success': True})


@socmed_bp.route('/api/scrape/tt', methods=['POST'])
@login_required
def api_scrape_tt():
    data = request.json
    user_id = data.get('username', '').strip()
    if not user_id:
        return jsonify({'success': False, 'error': 'User ID kosong'})
    try:
        url_endpoint = f'https://tiktok-best-experience.p.rapidapi.com/user/id/{user_id}'
        _rapid_key = os.environ.get('RAPIDAPI_KEY', '')
        if not _rapid_key:
            return jsonify({'success': False, 'error': 'RAPIDAPI_KEY belum diset di environment variable server.'})
        headers_rapid = {
            'X-RapidAPI-Key': _rapid_key,
            'X-RapidAPI-Host': 'tiktok-best-experience.p.rapidapi.com'
        }
        res_rapid = requests.get(url_endpoint, headers=headers_rapid, timeout=12)
        if res_rapid.status_code == 200:
            json_data = res_rapid.json()
            def find_key(keys, data_obj, default=None):
                if isinstance(data_obj, dict):
                    for k, v in data_obj.items():
                        if k in keys:
                            return v
                        found = find_key(keys, v, default)
                        if found is not default:
                            return found
                elif isinstance(data_obj, list):
                    for item in data_obj:
                        found = find_key(keys, item, default)
                        if found is not default:
                            return found
                return default
            followers_num = find_key(['follower_count', 'followerCount'], json_data, 0)
            posts_num = find_key(['aweme_count', 'videoCount'], json_data, 0)
            actual_username = find_key(['ins_id', 'uniqueId', 'unique_id'], json_data, f'user_{user_id}')
            full_name = find_key(['nickname', 'name'], json_data, f'ID: {user_id}')
            bio_text = find_key(['signature', 'desc'], json_data, '-')
            profile_pic = None
            avatar_obj = find_key(['avatar_300x300'], json_data)
            if isinstance(avatar_obj, dict):
                url_list = avatar_obj.get('url_list', [])
                if isinstance(url_list, list) and url_list:
                    profile_pic = url_list[0]
            if not profile_pic:
                profile_pic = find_key(['avatarLarger', 'avatarMedium'], json_data, f'https://ui-avatars.com/api/?name={actual_username}&background=random')
            return jsonify({
                'success': True,
                'is_fallback': False,
                'profile': {
                    'user_id': user_id,
                    'username': actual_username,
                    'profile_pic': profile_pic,
                    'bio': f'{full_name} | {bio_text}',
                    'followers': int(followers_num),
                    'total_posts': int(posts_num)
                }
            })
        return jsonify({'success': False, 'error': f'RapidAPI Error (Status: {res_rapid.status_code})'})
    except Exception as e:
        return jsonify({'success': False, 'error': f'Gagal memproses API: {str(e)}'})


@socmed_bp.route('/api/tt/auto_backup', methods=['POST'])
@login_required
def tt_auto_backup():
    accounts = load_json_file(TT_ACCOUNTS_FILE)
    if not accounts:
        return jsonify({'success': False, 'error': 'Database Kosong'})
    backup_dir = os.path.join('static', 'file_manager', 'Social Media Report', 'TT Report')
    os.makedirs(backup_dir, exist_ok=True)
    filename = f'TikTok_Report_AutoBackup_{datetime.now().strftime("%Y-%m-%d_%H%M")}.csv'
    filepath = os.path.join(backup_dir, filename)
    usernames = sorted(accounts.keys(), key=lambda x: accounts[x].get('profile', {}).get('followers', 0), reverse=True)
    with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(['No', 'Username', 'Nama Profil', 'Followers', 'Posts'])
        for i, u in enumerate(usernames):
            prof = accounts[u].get('profile', {})
            writer.writerow([i + 1, f"@{prof.get('username', u)}", prof.get('bio', '-').replace('\n', ' '), prof.get('followers', 0), prof.get('total_posts', 0)])
    return jsonify({'success': True, 'filename': filename})


@socmed_bp.route('/api/fb/accounts', methods=['GET', 'POST', 'DELETE'])
@login_required
def manage_fb_accounts():
    accounts = load_json_file(FB_ACCOUNTS_FILE)
    if request.method == 'GET':
        last_updated = '-'
        if os.path.exists(FB_ACCOUNTS_FILE):
            last_updated = datetime.fromtimestamp(os.path.getmtime(FB_ACCOUNTS_FILE)).strftime('%d %b %Y, %H:%M:%S')
        return jsonify({'accounts': accounts, 'last_updated': last_updated})
    if request.method == 'POST':
        data = request.json
        username = data.get('profile', {}).get('username')
        if username:
            existing = accounts.get(username, {})
            profile = data.get('profile', {})
            profile['posts_2025'] = existing.get('profile', {}).get('posts_2025', 0) if existing else profile.get('posts_2025', 0)
            profile['posts_2026'] = existing.get('profile', {}).get('posts_2026', {m: 0 for m in ['jan', 'feb', 'mar', 'apr', 'mei', 'jun', 'jul', 'agu', 'sep', 'okt', 'nov', 'des']}) if existing else profile.get('posts_2026', {m: 0 for m in ['jan', 'feb', 'mar', 'apr', 'mei', 'jun', 'jul', 'agu', 'sep', 'okt', 'nov', 'des']})
            profile['custom_name'] = existing.get('profile', {}).get('custom_name', '') if existing else profile.get('custom_name', '')
            total_2026 = sum(profile.get('posts_2026', {}).values())
            profile['total_posts'] = profile.get('posts_2025', 0) + total_2026
            accounts[username] = {'profile': profile}
            save_json_file(FB_ACCOUNTS_FILE, accounts)
        return jsonify({'success': True})
    username = request.json.get('username')
    if username in accounts:
        accounts.pop(username, None)
        save_json_file(FB_ACCOUNTS_FILE, accounts)
    return jsonify({'success': True})


@socmed_bp.route('/api/scrape/fb', methods=['POST'])
@login_required
def api_scrape_fb():
    data = request.json
    username = data.get('username', '').strip()
    if not username:
        return jsonify({'success': False, 'error': 'Username kosong'})
    try:
        followers_num = 0
        account_type = 'Personal Profile'
        bio_name = username
        profile_pic = f'https://ui-avatars.com/api/?name={username}&background=random'
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Accept-Language': 'id-ID,id;q=0.9,en-US;q=0.8,en;q=0.7'
        }
        res = requests.get(f'https://www.facebook.com/{username}/', headers=headers, timeout=12)
        soup = BeautifulSoup(res.text, 'html.parser')
        og_image = soup.find('meta', property='og:image')
        if og_image:
            profile_pic = og_image.get('content', profile_pic)
        og_title = soup.find('meta', property='og:title')
        if og_title:
            bio_name = og_title.get('content', username).split('|')[0].replace('- Home', '').replace('- Beranda', '').strip()
        else:
            title_tag = soup.find('title')
            if title_tag:
                bio_name = title_tag.get_text(strip=True).split('|')[0].replace('- Home', '').replace('- Beranda', '').strip()
        if bio_name.lower() == 'facebook':
            bio_name = username
        meta_desc = soup.find('meta', attrs={'name': 'description'}) or soup.find('meta', property='og:description')
        is_fanpage = False
        if meta_desc:
            content = meta_desc.get('content', '')
            if any(x in content.lower() for x in ['likes', 'menyukai', 'follower', 'pengikut']):
                is_fanpage = True
                account_type = 'Fanpage'
            f_match = re.search(r'([\d,KMB.rbjt\s]+)\s*(followers|pengikut)', content, re.IGNORECASE)
            if f_match:
                v = f_match.group(1).upper().replace(',', '.').replace(' ', '')
                if 'K' in v or 'RB' in v:
                    followers_num = int(float(v.replace('K', '').replace('RB', '')) * 1000)
                elif 'M' in v or 'JT' in v:
                    followers_num = int(float(v.replace('M', '').replace('JT', '')) * 1000000)
                else:
                    followers_num = int(float(v))
        final_followers = followers_num if is_fanpage else (followers_num if followers_num > 0 else 'Privat')
        return jsonify({'success': True, 'profile': {'username': username, 'profile_pic': profile_pic, 'bio': f'{bio_name} | {account_type}', 'followers': final_followers}})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@socmed_bp.route('/api/fb/save_posts', methods=['POST'])
@login_required
def api_fb_save_posts():
    data = request.json
    username = data.get('username')
    year = int(data.get('year', 2026))
    year_key = f'posts_{year}'
    year_2025 = data.get('posts_2025')
    months_data = data.get('months_data')
    custom_name = data.get('custom_name')
    if not username or not os.path.exists(FB_ACCOUNTS_FILE):
        return jsonify({'success': False, 'error': 'Parameter tidak valid'})
    accounts = load_json_file(FB_ACCOUNTS_FILE)
    if username in accounts:
        prof = accounts[username]['profile']
        if custom_name is not None:
            prof['custom_name'] = custom_name.strip()
        if year_2025 is not None:
            prof['posts_2025'] = int(year_2025)
        if months_data is not None:
            prof.setdefault(year_key, {})
            for m, val in months_data.items():
                prof[year_key][m] = int(val)
        tot = prof.get('posts_2025', 0)
        for k, v in prof.items():
            if k.startswith('posts_') and k != 'posts_2025':
                tot += sum(v.values())
        prof['total_posts'] = tot
        save_json_file(FB_ACCOUNTS_FILE, accounts)
        return jsonify({'success': True})
    return jsonify({'success': False, 'error': 'Akun tidak ditemukan'})


@socmed_bp.route('/api/fb/import_excel', methods=['POST'])
@login_required
def api_fb_import_excel():
    try:
        import openpyxl
    except ImportError:
        return jsonify({'success': False, 'error': 'Library openpyxl belum terinstall. Buka terminal: pip install openpyxl'})
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': 'Tidak ada file'})
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'error': 'File kosong'})
    try:
        wb = openpyxl.load_workbook(file, data_only=True)
        if os.path.exists(FB_ACCOUNTS_FILE):
            accounts = load_json_file(FB_ACCOUNTS_FILE)
        else:
            return jsonify({'success': False, 'error': 'Database kosong.'})
        updated_count = 0
        for sheet_name in wb.sheetnames:
            try:
                year = int(sheet_name)
            except ValueError:
                continue
            ws = wb[sheet_name]
            rows = list(ws.rows)
            if len(rows) < 2:
                continue
            headers = [str(h.value).strip().lower() if h.value else '' for h in rows[0]]
            try:
                idx_user = headers.index('username')
                months = ['jan', 'feb', 'mar', 'apr', 'mei', 'jun', 'jul', 'agu', 'sep', 'okt', 'nov', 'des']
                idx_months = {m: headers.index(m) for m in months}
            except ValueError:
                continue
            idx_prev = -1
            if year == 2026:
                try:
                    idx_prev = headers.index('post 2025')
                except ValueError:
                    pass
            for row in rows[1:]:
                if not row or len(row) <= idx_user or not row[idx_user].value:
                    continue
                raw_user = str(row[idx_user].value).replace('@', '').strip()
                if raw_user in accounts:
                    prof = accounts[raw_user]['profile']
                    if year == 2026 and idx_prev != -1:
                        try:
                            prof['posts_2025'] = int(row[idx_prev].value or 0)
                        except Exception:
                            pass
                    y_key = f'posts_{year}'
                    prof.setdefault(y_key, {})
                    for m_key, m_idx in idx_months.items():
                        try:
                            prof[y_key][m_key] = int(row[m_idx].value or 0)
                        except Exception:
                            prof[y_key][m_key] = 0
                    updated_count += 1
        for u in accounts:
            prof = accounts[u]['profile']
            tot = prof.get('posts_2025', 0)
            for k, v in prof.items():
                if k.startswith('posts_') and k != 'posts_2025':
                    tot += sum(v.values())
            prof['total_posts'] = tot
        save_json_file(FB_ACCOUNTS_FILE, accounts)
        return jsonify({'success': True, 'updated': updated_count})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@socmed_bp.route('/api/<platform>/import_excel', methods=['POST'])
@login_required
def api_ig_tt_import_excel(platform):
    if platform not in ['ig', 'tt']:
        return jsonify({'error': 'Platform tidak valid'})
    target_file = IG_ACCOUNTS_FILE if platform == 'ig' else TT_ACCOUNTS_FILE
    try:
        import openpyxl
    except ImportError:
        return jsonify({'success': False, 'error': 'Library openpyxl belum terinstall.'})
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': 'Tidak ada file'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'error': 'File kosong'}), 400
    try:
        wb = openpyxl.load_workbook(file, data_only=True)
        if os.path.exists(target_file):
            accounts = load_json_file(target_file)
        else:
            return jsonify({'success': False, 'error': 'Database diler kosong. Harap add akun dulu di Dasbor.'})
        updated_count = 0
        valid_years = []
        for sheet_name in wb.sheetnames:
            try:
                valid_years.append(int(sheet_name.strip()))
            except ValueError:
                continue
        valid_years.sort()
        if not valid_years:
            return jsonify({'success': False, 'error': 'Tidak ditemukan nama sheet berformat tahun (contoh: 2026, 2027)'})
        for year in valid_years:
            ws = wb[str(year)]
            rows = list(ws.rows)
            if len(rows) < 2:
                continue
            headers = [str(h.value).strip().lower() if h.value else '' for h in rows[0]]
            try:
                idx_user = headers.index('username')
                months = ['jan', 'feb', 'mar', 'apr', 'mei', 'jun', 'jul', 'agu', 'sep', 'okt', 'nov', 'des']
                idx_months = {m: headers.index(m) for m in months if m in headers}
            except ValueError:
                continue
            baseline_col_name = f'post {year - 1}'
            idx_baseline = -1
            try:
                idx_baseline = headers.index(baseline_col_name)
            except ValueError:
                pass
            for row in rows[1:]:
                if not row or len(row) <= idx_user or not row[idx_user].value:
                    continue
                raw_user = str(row[idx_user].value).replace('@', '').strip().lower()
                prof = None
                if platform == 'ig':
                    if raw_user in accounts and 'profile' in accounts[raw_user]:
                        prof = accounts[raw_user]['profile']
                else:
                    for key_id, target_obj in accounts.items():
                        if isinstance(target_obj, dict) and 'profile' in target_obj and str(target_obj['profile'].get('username', '')).lower() == raw_user:
                            prof = target_obj['profile']
                            break
                if prof is None:
                    continue
                if idx_baseline != -1:
                    try:
                        prof[f'posts_base_{year}'] = int(row[idx_baseline].value or 0)
                    except Exception:
                        pass
                y_posts_key = f'posts_{year}'
                y_foll_key = f'followers_{year}'
                prof.setdefault(y_posts_key, {})
                prof.setdefault(y_foll_key, {})
                for m_key, m_idx in idx_months.items():
                    try:
                        val = int(row[m_idx].value or 0)
                    except Exception:
                        val = 0
                    prof[y_posts_key][m_key] = val
                    prof[y_foll_key].setdefault(m_key, 0)
                updated_count += 1
                if idx_baseline != -1:
                    try:
                        total_posts_this_year = sum(prof[y_posts_key].values())
                        prof['posts_base'] = prof[f'posts_base_{year}'] + total_posts_this_year
                    except Exception:
                        pass
        save_json_file(target_file, accounts)
        return jsonify({'success': True, 'updated': updated_count, 'years_processed': valid_years})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@socmed_bp.route('/api/socmed/stats', methods=['GET'])
@login_required
def api_socmed_stats():
    def process_accounts(file_path):
        data = {'labels': [], 'names': [], 'pics': [], 'accounts_data': []}
        last_updated = 'Belum ada data'
        if os.path.exists(file_path):
            last_updated = datetime.fromtimestamp(os.path.getmtime(file_path)).strftime('%d %b %Y, %H:%M:%S')
            accounts = load_json_file(file_path)
            def get_sort_value(x):
                val = x.get('profile', {}).get('followers', 0)
                return val if isinstance(val, (int, float)) else 0
            sorted_accs = sorted(accounts.values(), key=get_sort_value, reverse=True)
            for acc in sorted_accs:
                prof = acc.get('profile', {})
                username = prof.get('username', '')
                bio_raw = prof.get('bio', username)
                full_name = prof.get('custom_name') or (bio_raw.split('|')[0].strip() if '|' in bio_raw else bio_raw)
                data['labels'].append(f"@{username}")
                data['names'].append(full_name)
                data['pics'].append(prof.get('profile_pic', f'https://ui-avatars.com/api/?name={username}&background=random'))
                data['accounts_data'].append(prof)
        return data, last_updated
    ig_data, ig_updated = process_accounts(IG_ACCOUNTS_FILE)
    tt_data, tt_updated = process_accounts(TT_ACCOUNTS_FILE)
    fb_data, fb_updated = process_accounts(FB_ACCOUNTS_FILE)
    return jsonify({'instagram': {**ig_data, 'last_updated': ig_updated}, 'tiktok': {**tt_data, 'last_updated': tt_updated}, 'facebook': {**fb_data, 'last_updated': fb_updated}})


@socmed_bp.route('/api/gmaps/accounts', methods=['GET', 'POST', 'DELETE'])
@login_required
def manage_gmaps_accounts():
    accounts = load_json_file(GMAPS_ACCOUNTS_FILE)
    if request.method == 'GET':
        last_updated = '-'
        if os.path.exists(GMAPS_ACCOUNTS_FILE):
            last_updated = datetime.fromtimestamp(os.path.getmtime(GMAPS_ACCOUNTS_FILE)).strftime('%d %b %Y, %H:%M:%S')
        
        month_filter = request.args.get('month')
        if month_filter:
            reviews_db = load_json_file(GMAPS_REVIEWS_FILE)
            for b_id, acc in accounts.items():
                prof = acc.get('profile', {}) if isinstance(acc, dict) else {}
                b_reviews = reviews_db.get(b_id, [])
                
                # Reset metrics
                m_reviews = 0
                m_stars = {1:0, 2:0, 3:0, 4:0, 5:0}
                
                for r in b_reviews:
                    if r.get('date', '').startswith(month_filter):
                        m_reviews += 1
                        rtg = int(r.get('rating', 0))
                        if rtg in m_stars:
                            m_stars[rtg] += 1
                
                prof['reviews'] = m_reviews
                prof['star_1'] = m_stars[1]
                prof['star_2'] = m_stars[2]
                prof['star_3'] = m_stars[3]
                prof['star_4'] = m_stars[4]
                prof['star_5'] = m_stars[5]
                # Note: We leave rating as is (overall rating)
        
        return jsonify({'accounts': accounts, 'last_updated': last_updated})
    if request.method == 'POST':
        data = request.json
        key_id = data.get('profile', {}).get('business_id')
        if key_id:
            accounts[key_id] = data
            save_json_file(GMAPS_ACCOUNTS_FILE, accounts)
        return jsonify({'success': True})
    key_id = request.json.get('business_id')
    if key_id in accounts:
        accounts.pop(key_id, None)
        save_json_file(GMAPS_ACCOUNTS_FILE, accounts)
        
        # Juga hapus dari reviews
        reviews_data = load_json_file(GMAPS_REVIEWS_FILE)
        if key_id in reviews_data:
            reviews_data.pop(key_id, None)
            save_json_file(GMAPS_REVIEWS_FILE, reviews_data)
            
    return jsonify({'success': True})

import threading
import time

MASS_UPDATE_STATUS = {
    'running': False,
    'progress': 0,
    'total': 0,
    'message': '',
    'log': ''
}

def bg_mass_update_gmaps():
    global MASS_UPDATE_STATUS
    MASS_UPDATE_STATUS['running'] = True
    MASS_UPDATE_STATUS['message'] = 'Memulai proses...'
    
    try:
        accounts = load_json_file(GMAPS_ACCOUNTS_FILE)
        if not accounts:
            MASS_UPDATE_STATUS['message'] = 'Database kosong.'
            return
            
        apify_token = os.environ.get('APIFY_TOKEN', '')
        if not apify_token:
            MASS_UPDATE_STATUS['message'] = 'APIFY_TOKEN tidak ditemukan di .env'
            return
            
        bids = list(accounts.keys())
        MASS_UPDATE_STATUS['total'] = len(bids)
        MASS_UPDATE_STATUS['progress'] = 0
        
        MASS_UPDATE_STATUS['message'] = 'Menarik data profil massal via Apify...'
        
        start_urls = []
        for bid in bids:
            prof = accounts[bid].get('profile', {})
            link = prof.get('link')
            
            if link and link.startswith('http'):
                start_urls.append({"url": link})
            elif bid.startswith('ChIJ'):
                start_urls.append({"url": f"https://www.google.com/maps/place/?q=place_id:{bid}"})
            else:
                query = prof.get('name', bid)
                start_urls.append({"url": f"https://www.google.com/maps/search/?api=1&query={query}"})
            
        try:
            client = ApifyClient(apify_token)
            run_input = {
                "startUrls": start_urls,
                "maxCrawledPlacesPerSearch": 1,
                "maxReviews": 0,
                "maxImages": 0,
                "language": "id"
            }
            
            run_info = client.actor("compass/crawler-google-places").start(run_input=run_input)
            run_id = run_info.get('id') if isinstance(run_info, dict) else getattr(run_info, 'id', None)
            
            MASS_UPDATE_STATUS['message'] = 'Menunggu Apify...'
            MASS_UPDATE_STATUS['log'] = 'Memulai container Apify...\n'
            
            import time
            while True:
                info = client.run(run_id).get()
                status = info.get('status') if isinstance(info, dict) else getattr(info, 'status', None)
                
                try:
                    log_bytes = client.run(run_id).log().get()
                    if log_bytes:
                        log_str = log_bytes.decode('utf-8') if isinstance(log_bytes, bytes) else str(log_bytes)
                        lines = log_str.strip().split('\n')
                        MASS_UPDATE_STATUS['log'] = '\n'.join(lines[-10:])
                except:
                    pass
                    
                if status in ['SUCCEEDED', 'FAILED', 'ABORTED']:
                    if status != 'SUCCEEDED':
                        MASS_UPDATE_STATUS['message'] = f'Apify Error: {status}'
                        MASS_UPDATE_STATUS['running'] = False
                        return
                    break
                time.sleep(3)
                
            MASS_UPDATE_STATUS['log'] += 'Proses Apify selesai. Mengekstrak data...\n'
            dataset_id = getattr(info, 'default_dataset_id', None) or (info.get('defaultDatasetId') if isinstance(info, dict) else None)
            dataset_items = list(client.dataset(dataset_id).iterate_items())
            
            # Map results to accounts
            for idx, d in enumerate(dataset_items):
                MASS_UPDATE_STATUS['progress'] = idx + 1
                title = d.get('title', '').lower()
                
                # Cari akun yang cocok dengan hasil Apify
                for bid, acc in accounts.items():
                    prof = acc.get('profile', {})
                    acc_name = prof.get('name', bid).lower()
                    
                    if acc_name in title or title in acc_name:
                        prof['rating'] = float(d.get('totalScore', prof.get('rating', 0.0)))
                        prof['reviews'] = int(d.get('reviewsCount', prof.get('reviews', 0)))
                        
                        loc = d.get('location', {})
                        if loc:
                            prof['latitude'] = loc.get('lat', prof.get('latitude', '-'))
                            prof['longitude'] = loc.get('lng', prof.get('longitude', '-'))
                            
                        # Update accounts
                        accounts[bid]['profile'] = prof
                        break
                        
            save_json_file(GMAPS_ACCOUNTS_FILE, accounts)
            MASS_UPDATE_STATUS['message'] = 'Selesai'
            
        except Exception as e:
            MASS_UPDATE_STATUS['message'] = f'Error API: {str(e)}'
        
    except Exception as e:
        MASS_UPDATE_STATUS['message'] = f'Error: {str(e)}'
        
    finally:
        MASS_UPDATE_STATUS['running'] = False


@socmed_bp.route('/api/gmaps/mass_update', methods=['POST'])
@login_required
def api_gmaps_mass_update():
    global MASS_UPDATE_STATUS
    if MASS_UPDATE_STATUS['running']:
        return jsonify({'success': False, 'error': 'Proses update masih berjalan'})
        
    thread = threading.Thread(target=bg_mass_update_gmaps)
    thread.daemon = True
    thread.start()
    
    return jsonify({'success': True, 'message': 'Update massal dimulai'})

@socmed_bp.route('/api/gmaps/mass_update/status', methods=['GET'])
@login_required
def api_gmaps_mass_update_status():
    return jsonify(MASS_UPDATE_STATUS)

GRAB_REVIEWS_STATUS = {
    'running': False,
    'progress': 0,
    'total': 0,
    'message': '',
    'log': ''
}

def bg_grab_reviews(business_id):
    global GRAB_REVIEWS_STATUS
    GRAB_REVIEWS_STATUS['running'] = True
    GRAB_REVIEWS_STATUS['message'] = 'Memulai proses...'
    GRAB_REVIEWS_STATUS['progress'] = 0
    GRAB_REVIEWS_STATUS['total'] = 0
    
    try:
        accounts = load_json_file(GMAPS_ACCOUNTS_FILE)
        if business_id not in accounts:
            GRAB_REVIEWS_STATUS['message'] = 'Bisnis tidak ditemukan di database.'
            return
            
        reviews_db = load_json_file(GMAPS_REVIEWS_FILE)
        business_reviews = reviews_db.get(business_id, [])
        
        def _clean(t):
            if not t: return ""
            return "".join(c for c in str(t).lower() if c.isalnum())
            
        existing_ids = set()
        existing_sigs = set()
        for r in business_reviews:
            r_id = r.get('review_id')
            if r_id:
                existing_ids.add(r_id)
            sig = f"{_clean(r.get('reviewer'))}||{_clean(r.get('text'))}"
            existing_sigs.add(sig)
        
        apify_token = os.environ.get('APIFY_TOKEN', '')
        if not apify_token:
            GRAB_REVIEWS_STATUS['message'] = 'APIFY_TOKEN tidak ditemukan di .env'
            return
            
        prof = accounts[business_id].get('profile', {})
        
        # Hitung dynamic maxReviews
        google_total = int(prof.get('reviews', 0))
        db_total = len(business_reviews)
        expected_new = max(0, google_total - db_total)
        max_reviews = max(50, expected_new + 20) # Minimal 50, buffer 20
        
        GRAB_REVIEWS_STATUS['message'] = f'Menarik ulasan via Apify (Max: {max_reviews})...'
        
        new_reviews = []
        try:
            client = ApifyClient(apify_token)
            run_input = {
                "startUrls": [{"url": f"https://www.google.com/maps/place/?q=place_id:{business_id}"}],
                "maxCrawledPlacesPerSearch": 1,
                "maxReviews": max_reviews,
                "reviewsSort": "newest",
                "language": "id"
            }
            
            run_info = client.actor("compass/crawler-google-places").start(run_input=run_input)
            run_id = run_info.get('id') if isinstance(run_info, dict) else getattr(run_info, 'id', None)
            
            GRAB_REVIEWS_STATUS['message'] = 'Menunggu Apify...'
            GRAB_REVIEWS_STATUS['log'] = 'Memulai container Apify...\n'
            
            import time
            while True:
                info = client.run(run_id).get()
                status = info.get('status') if isinstance(info, dict) else getattr(info, 'status', None)
                
                try:
                    log_bytes = client.run(run_id).log().get()
                    if log_bytes:
                        log_str = log_bytes.decode('utf-8') if isinstance(log_bytes, bytes) else str(log_bytes)
                        lines = log_str.strip().split('\n')
                        GRAB_REVIEWS_STATUS['log'] = '\n'.join(lines[-10:])
                except:
                    pass
                    
                if status in ['SUCCEEDED', 'FAILED', 'ABORTED']:
                    if status != 'SUCCEEDED':
                        GRAB_REVIEWS_STATUS['message'] = f'Apify Error: {status}'
                        GRAB_REVIEWS_STATUS['running'] = False
                        return
                    break
                time.sleep(3)
                
            GRAB_REVIEWS_STATUS['log'] += 'Proses Apify selesai. Mengekstrak data...\n'
            dataset_id = getattr(info, 'default_dataset_id', None) or (info.get('defaultDatasetId') if isinstance(info, dict) else None)
            dataset_items = list(client.dataset(dataset_id).iterate_items())
            
            if not dataset_items:
                GRAB_REVIEWS_STATUS['message'] = 'Error: Apify tidak menemukan data (Kemungkinan ID/URL tidak valid atau struktur halaman Google berubah)'
                return
                
            place_data = dataset_items[0]
            data_reviews = place_data.get('reviews', [])
            
            new_reviews = []
            stop_early = False
            for r in data_reviews:
                r_id = r.get('reviewId', r.get('id'))
                author = r.get('reviewerName', r.get('name', 'Anonim'))
                text = r.get('text', '')
                sig = f"{_clean(author)}||{_clean(text)}"
                
                if (r_id and r_id in existing_ids) or (sig in existing_sigs):
                    stop_early = True
                    break
                
                parsed_r = {
                    'review_id': r_id,
                    'reviewer': author,
                    'date': r.get('publishDate', r.get('publishedAtDate', '')),
                    'rating': int(r.get('stars') or r.get('rating') or 0),
                    'text': text,
                    'source': 'Google Maps'
                }
                new_reviews.append(parsed_r)
                GRAB_REVIEWS_STATUS['progress'] = len(new_reviews)
                
        except Exception as e:
            GRAB_REVIEWS_STATUS['message'] = f'Error API: {str(e)}'
            
        if new_reviews:
            if business_id in reviews_db:
                merged = {}
                for r in reviews_db[business_id]:
                    rid = r.get('review_id')
                    sig = f"{_clean(r.get('reviewer'))}||{_clean(r.get('text'))}"
                    merged[rid if rid else sig] = r
                for r in reversed(new_reviews):
                    rid = r.get('review_id')
                    sig = f"{_clean(r.get('reviewer'))}||{_clean(r.get('text'))}"
                    merged[rid if rid else sig] = r
                final_list = list(merged.values())
                final_list.sort(key=lambda x: x.get('date', ''), reverse=True)
                reviews_db[business_id] = final_list
            else:
                reviews_db[business_id] = new_reviews
                
            save_json_file(GMAPS_REVIEWS_FILE, reviews_db)
            GRAB_REVIEWS_STATUS['message'] = 'Selesai'
        else:
            if 'Error API' not in GRAB_REVIEWS_STATUS['message'] and 'API Error' not in GRAB_REVIEWS_STATUS['message'] and 'Endpoint' not in GRAB_REVIEWS_STATUS['message']:
                GRAB_REVIEWS_STATUS['message'] = 'Selesai. Tidak ada ulasan baru.'
        
        if 'Error API' not in GRAB_REVIEWS_STATUS['message'] and 'API Error' not in GRAB_REVIEWS_STATUS['message'] and 'Endpoint' not in GRAB_REVIEWS_STATUS['message'] and 'Error:' not in GRAB_REVIEWS_STATUS['message']:
            prof = accounts[business_id]['profile']
            prof['last_reviews_update'] = datetime.now().strftime('%d %b %Y, %H:%M:%S')
            
            # Kalkulasi ulang bintang berdasarkan review lokal
            all_revs = reviews_db.get(business_id, [])
            star_counts = {1:0, 2:0, 3:0, 4:0, 5:0}
            for rv in all_revs:
                r_rating = rv.get('rating', 0)
                if r_rating in star_counts:
                    star_counts[r_rating] += 1
            prof['star_1'] = star_counts[1]
            prof['star_2'] = star_counts[2]
            prof['star_3'] = star_counts[3]
            prof['star_4'] = star_counts[4]
            prof['star_5'] = star_counts[5]
            
            save_json_file(GMAPS_ACCOUNTS_FILE, accounts)
            
    except Exception as e:
        GRAB_REVIEWS_STATUS['message'] = f'Error: {str(e)}'
        
    finally:
        GRAB_REVIEWS_STATUS['running'] = False


@socmed_bp.route('/api/gmaps/grab_reviews', methods=['POST'])
@login_required
def api_gmaps_grab_reviews():
    global GRAB_REVIEWS_STATUS
    if GRAB_REVIEWS_STATUS['running']:
        return jsonify({'success': False, 'error': 'Proses grab masih berjalan'})
        
    business_id = request.json.get('business_id')
    if not business_id:
         return jsonify({'success': False, 'error': 'Business ID kosong'})
         
    thread = threading.Thread(target=bg_grab_reviews, args=(business_id,))
    thread.daemon = True
    thread.start()
    
    return jsonify({'success': True, 'message': 'Grab dimulai'})

@socmed_bp.route('/api/gmaps/grab_reviews/status', methods=['GET'])
@login_required
def api_gmaps_grab_reviews_status():
    return jsonify(GRAB_REVIEWS_STATUS)




@socmed_bp.route('/api/scrape/gmaps', methods=['POST'])
@login_required
def api_scrape_gmaps():
    data = request.json
    business_id = data.get('business_id', '').strip()
    if not business_id:
        return jsonify({'success': False, 'error': 'Business ID kosong'})
    
    try:
        apify_token = os.environ.get('APIFY_TOKEN', '')
        if not apify_token:
            return jsonify({'success': False, 'error': 'APIFY_TOKEN belum diset di environment variable server.'})
        
        client = ApifyClient(apify_token)
        
        start_urls = []
        if business_id.startswith('http'):
            start_urls.append({"url": business_id})
        elif business_id.startswith('ChIJ'):
            start_urls.append({"url": f"https://www.google.com/maps/place/?q=place_id:{business_id}"})
        else:
            start_urls.append({"url": f"https://www.google.com/maps/search/?api=1&query={business_id}"})
            
        run_input = {
            "startUrls": start_urls,
            "maxCrawledPlacesPerSearch": 1,
            "maxReviews": 0,
            "maxImages": 0,
            "language": "id"
        }
        
        run = client.actor("compass/crawler-google-places").call(run_input=run_input)
        dataset_id = run.get('defaultDatasetId') if isinstance(run, dict) else getattr(run, 'default_dataset_id', None)
        dataset_items = list(client.dataset(dataset_id).iterate_items())
        
        if len(dataset_items) > 0:
            data_obj = dataset_items[0]
            
            name = data_obj.get('title', business_id)
            rating = data_obj.get('totalScore', 0.0)
            reviews = data_obj.get('reviewsCount', 0)
            
            loc = data_obj.get('location', {})
            latitude = loc.get('lat', '-')
            longitude = loc.get('lng', '-')
            
            # Ambil distribusi bintang jika Apify menyediakannya
            dist = data_obj.get('reviewsDistribution', {})
            star_1 = dist.get('oneStar', 0)
            star_2 = dist.get('twoStar', 0)
            star_3 = dist.get('threeStar', 0)
            star_4 = dist.get('fourStar', 0)
            star_5 = dist.get('fiveStar', 0)
            
            actual_bid = data_obj.get('placeId', business_id)

            return jsonify({
                'success': True,
                'is_fallback': False,
                'profile': {
                    'business_id': actual_bid,
                    'name': name,
                    'rating': float(rating),
                    'reviews': int(reviews),
                    'star_1': int(star_1),
                    'star_2': int(star_2),
                    'star_3': int(star_3),
                    'star_4': int(star_4),
                    'star_5': int(star_5),
                    'latitude': latitude,
                    'longitude': longitude,
                    'link': data_obj.get('url', '')
                }
            })
        else:
            return jsonify({'success': False, 'error': 'Tempat tidak ditemukan di Google Maps.'})
            
    except Exception as e:
        return jsonify({'success': False, 'error': f'Gagal memproses API: {str(e)}'})

@socmed_bp.route('/api/gmaps/auto_backup', methods=['POST'])
@login_required
def gmaps_auto_backup():
    accounts = load_json_file(GMAPS_ACCOUNTS_FILE)
    if not accounts:
        return jsonify({'success': False, 'error': 'Database Kosong'})
    backup_dir = os.path.join('static', 'file_manager', 'Google Maps Report')
    os.makedirs(backup_dir, exist_ok=True)
    filename = f'GMaps_Report_AutoBackup_{datetime.now().strftime("%Y-%m-%d_%H%M")}.csv'
    filepath = os.path.join(backup_dir, filename)
    bids = sorted(accounts.keys(), key=lambda x: accounts[x].get('profile', {}).get('reviews', 0), reverse=True)
    with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(['No', 'Business Name', 'Rating', 'Total Reviews', 'Star 1', 'Star 2', 'Star 3', 'Star 4', 'Star 5', 'Koordinat'])
        for i, u in enumerate(bids):
            prof = accounts[u].get('profile', {})
            coord = f"{prof.get('latitude', '-')}, {prof.get('longitude', '-')}"
            writer.writerow([
                i + 1,
                prof.get('name', u),
                prof.get('rating', 0),
                prof.get('reviews', 0),
                prof.get('star_1', 0),
                prof.get('star_2', 0),
                prof.get('star_3', 0),
                prof.get('star_4', 0),
                prof.get('star_5', 0),
                coord
            ])
    return jsonify({'success': True, 'filename': filename})


@socmed_bp.route('/api/gmaps/reviews/<business_id>', methods=['GET'])
@login_required
def get_gmaps_reviews(business_id):
    year = request.args.get('year')
    all_reviews = load_json_file(GMAPS_REVIEWS_FILE)
    
    b_reviews = all_reviews.get(business_id, [])
    
    # Sort reviews by date descending
    b_reviews.sort(key=lambda x: x.get('date', ''), reverse=True)
    
    # Monthly Stats if year is requested
    monthly_stats = None
    historical_total = 0
    if year:
        # Initialize 12 months
        monthly_stats = {
            m: {'rating': 0, 'reviews': 0, 'star_1': 0, 'star_2': 0, 'star_3': 0, 'star_4': 0, 'star_5': 0, 'total_score': 0}
            for m in range(1, 13)
        }
        
        for r in b_reviews:
            date_str = r.get('date', '')
            if date_str:
                if date_str < f"{year}-01-01":
                    historical_total += 1
                elif date_str.startswith(year):
                    try:
                        # Format is usually YYYY-MM-DD HH:MM:SS
                        m = int(date_str.split('-')[1])
                        rtg = r.get('rating', 0)
                        monthly_stats[m]['reviews'] += 1
                        monthly_stats[m]['total_score'] += rtg
                        if rtg == 1: monthly_stats[m]['star_1'] += 1
                        elif rtg == 2: monthly_stats[m]['star_2'] += 1
                        elif rtg == 3: monthly_stats[m]['star_3'] += 1
                        elif rtg == 4: monthly_stats[m]['star_4'] += 1
                        elif rtg == 5: monthly_stats[m]['star_5'] += 1
                    except:
                        pass
        
        # Calculate average ratings
        for m in range(1, 13):
            if monthly_stats[m]['reviews'] > 0:
                monthly_stats[m]['rating'] = round(monthly_stats[m]['total_score'] / monthly_stats[m]['reviews'], 2)
            del monthly_stats[m]['total_score']
            
    # Process adjustments removed as per user request
            
    # Get Google's official total reviews count
    all_accounts = load_json_file(GMAPS_ACCOUNTS_FILE)
    b_acc = all_accounts.get(business_id, {})
    b_prof = b_acc.get('profile', {}) if isinstance(b_acc, dict) else {}
    google_total = b_prof.get('reviews', len(b_reviews))
    last_update = b_prof.get('last_reviews_update', '-')
            
    return jsonify({
        'success': True,
        'business_id': business_id,
        'total': len(b_reviews),
        'google_total': google_total,
        'last_update': last_update,
        'reviews': b_reviews,
        'monthly_stats': monthly_stats,
        'historical_total': historical_total
    })

@socmed_bp.route('/api/socmed/force_monthly_delta', methods=['POST'])
@login_required
def api_force_monthly_delta():
    from ..cron import calculate_monthly_delta
    try:
        calculate_monthly_delta()
        return jsonify({'success': True, 'message': 'Kalkulasi data bulanan berhasil dieksekusi.'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})
